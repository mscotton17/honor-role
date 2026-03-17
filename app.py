import io
import pandas as pd
import streamlit as st
from docx import Document
from docx.shared import Pt

st.set_page_config(page_title="Honor Roll Protocol", page_icon="🎓")
st.title("Honor Roll Protocol")
st.write("Upload a student CSV file to generate a formatted Honor Roll Word document.")

uploaded_file = st.file_uploader("Choose a CSV file", type=["csv"])

if uploaded_file is not None:
    try:
        df = pd.read_csv(uploaded_file, skiprows=2)
    except Exception as e:
        st.error(f"Failed to read CSV: {e}")
        st.stop()

    required_cols = {"FIRST_NAME", "LAST_NAME", "PHYSICAL_CITY", "PHYSICAL_STATE"}
    missing = required_cols - set(df.columns)
    if missing:
        st.error(f"CSV is missing required columns: {', '.join(missing)}")
        st.stop()

    # Drop rows missing name or city/state data
    df = df.dropna(subset=["FIRST_NAME", "LAST_NAME", "PHYSICAL_CITY", "PHYSICAL_STATE"])

    # Strip whitespace
    for col in ["FIRST_NAME", "LAST_NAME", "PHYSICAL_CITY", "PHYSICAL_STATE"]:
        df[col] = df[col].astype(str).str.strip()

    # Remove rows where key fields became empty strings after stripping
    df = df[
        df["FIRST_NAME"].ne("") &
        df["LAST_NAME"].ne("") &
        df["PHYSICAL_CITY"].ne("") &
        df["PHYSICAL_STATE"].ne("")
    ]

    if df.empty:
        st.warning("No valid student records found after cleaning the data.")
        st.stop()

    df["Full Name"] = df["FIRST_NAME"] + " " + df["LAST_NAME"]
    df = df.sort_values(["PHYSICAL_CITY", "PHYSICAL_STATE", "LAST_NAME"])

    st.success(f"Loaded {len(df)} student records.")
    st.dataframe(
        df[["Full Name", "PHYSICAL_CITY", "PHYSICAL_STATE"]].reset_index(drop=True),
        use_container_width=True,
    )

    # Build Word document
    doc = Document()

    current_location = None
    for _, row in df.iterrows():
        location = f"{row['PHYSICAL_CITY']}, {row['PHYSICAL_STATE']}"
        if location != current_location:
            heading = doc.add_heading(location, level=2)
            for run in heading.runs:
                run.bold = True
            current_location = location
        doc.add_paragraph(row["Full Name"], style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    st.download_button(
        label="Download Honor_Roll_Final.docx",
        data=buf,
        file_name="Honor_Roll_Final.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
